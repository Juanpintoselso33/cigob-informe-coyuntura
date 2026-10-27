"""ADR-0385: siniestralidad laboral registrada en el sistema de riesgos del trabajo.

Todo sin red: el Cuadro 1 sale de fixtures con el texto REAL que extrae
pdfplumber de dos ediciones del informe provisorio de la SRT (enero-junio de
2026 y enero-marzo de 2025, que tienen filas con y sin rótulo en la misma línea).
"""
import io
import json
import shutil
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(ROOT))
import config  # noqa: E402
import gestion  # noqa: E402
import itcg  # noqa: E402
import parametrica  # noqa: E402
import procedencia_anclas  # noqa: E402

FIX = ROOT / "tests" / "fixtures"
TEXTO_2026 = (FIX / "srt_cuadro1_2026_ene_jun.txt").read_text(encoding="utf-8")
TEXTO_2025 = (FIX / "srt_cuadro1_2025_ene_mar.txt").read_text(encoding="utf-8")
STORE_REAL = ROOT / "data" / "gestion" / "srt_siniestralidad.json"


# ── Parser del Cuadro 1 ──────────────────────────────────────────────────────

def test_lee_el_cuadro_1_de_enero_a_junio_de_2026():
    ed = gestion.srt_parsear_provisorio(TEXTO_2026)
    assert (ed["anio"], ed["hasta_mes"]) == (2026, 6)
    assert ed["periodo_txt"] == "enero a junio de 2026"
    assert ed["elaboracion"] == "septiembre de 2026"
    a26, a25, a24 = ed["anios"]["2026"], ed["anios"]["2025"], ed["anios"]["2024"]
    # Total sistema, casos AT y EP con días de baja o secuelas (sin in itinere).
    assert a26["casos"] == {"unidades": 146899, "casas": 2167, "total": 149066}
    assert a25["casos"]["total"] == 159671
    assert a24["casos"]["total"] == 151352
    assert a26["cubiertos"]["total"] == 9991247
    assert a25["cubiertos"]["total"] == 10167202


def test_lee_el_cuadro_1_de_otra_edicion_con_otro_periodo():
    ed = gestion.srt_parsear_provisorio(TEXTO_2025)
    assert (ed["anio"], ed["hasta_mes"]) == (2025, 3)
    assert ed["periodo_txt"] == "enero a marzo de 2025"
    assert ed["anios"]["2025"]["casos"]["total"] == 77373
    assert ed["anios"]["2024"]["casos"]["total"] == 76547
    assert ed["anios"]["2025"]["cubiertos"]["total"] == 10139577


def test_las_personas_cubiertas_pueden_diferir_en_uno_pero_los_casos_no():
    """Enero-junio de 2024: 9.716.834 + 532.197 = 10.249.031 y el informe
    imprime 10.249.030 (son promedios redondeados). No es un error de lectura."""
    assert 9716834 + 532197 != 10249030
    gestion.srt_parsear_provisorio(TEXTO_2026)        # no levanta
    roto = TEXTO_2026.replace("149.600 1.752 151.352", "149.600 1.752 151.353")
    with pytest.raises(ValueError, match="no suman el total"):
        gestion.srt_parsear_provisorio(roto)


def test_falla_en_voz_alta_si_un_total_no_cierra():
    roto = TEXTO_2026.replace("146.899 2.167 149.066", "146.899 2.167 149.166")
    with pytest.raises(ValueError):
        gestion.srt_parsear_provisorio(roto)


def test_falla_si_el_cuadro_2_imprime_otra_variacion():
    """Protege contra leer la fila equivocada: el informe imprime -6,6 % y la
    cuenta propia tiene que coincidir."""
    roto = TEXTO_2026.replace("* -6,9% 13,8% -6,6%", "* -6,9% 13,8% -2,2%")
    assert roto != TEXTO_2026
    with pytest.raises(ValueError, match="Cuadro 2"):
        gestion.srt_parsear_provisorio(roto)


def test_falla_si_faltan_las_filas_o_el_titulo():
    sin_filas = "\n".join(l for l in TEXTO_2026.splitlines()
                          if not l.strip().startswith(("9.716.834", "249.019", "60.917", "149.600")))
    with pytest.raises(ValueError):
        gestion.srt_parsear_provisorio(sin_filas)
    with pytest.raises(ValueError, match="período"):
        gestion.srt_parsear_provisorio(TEXTO_2026.replace("Enero a junio", "Primer semestre"))
    with pytest.raises(ValueError):
        gestion.srt_parsear_provisorio("el PDF cambió de diseño")


# ── La cuenta ────────────────────────────────────────────────────────────────

def _store_con(*textos):
    store = {"anuales": {}, "ediciones": {}}
    for t in textos:
        gestion.srt_incorporar_edicion(store, gestion.srt_parsear_provisorio(t), leida="2026-10-27")
    return store


def test_el_valor_es_la_variacion_del_indice_de_incidencia():
    p = gestion.siniestralidad_puntos(_store_con(TEXTO_2026))[-1]
    # (149.066 / 9.991.247) / (159.671 / 10.167.202) − 1
    esperado = ((149066 / 9991247) / (159671 / 10167202) - 1) * 100
    assert p["valor"] == round(esperado, 1) == -5.0
    assert p["fecha"] == "2026-06-01" and p["tipo"] == "acumulado"
    assert (p["indice_actual"], p["indice_previo"]) == (14.92, 15.7)


def test_el_indice_no_es_la_variacion_de_los_casos():
    """Los casos bajan 6,6 % pero los cubiertos también bajan: lo que mide el
    indicador es la incidencia, no el conteo."""
    p = gestion.siniestralidad_puntos(_store_con(TEXTO_2026))[-1]
    assert p["valor"] != round((149066 / 159671 - 1) * 100, 1)


def test_los_puntos_anuales_y_acumulados_salen_ordenados():
    store = _store_con(TEXTO_2025, TEXTO_2026)
    store["anuales"] = {"2024": {"casos": 326004, "incidencia_mil": 32.0},
                        "2025": {"casos": 332708, "incidencia_mil": 32.8}}
    puntos = gestion.siniestralidad_puntos(store)
    assert [p["fecha"] for p in puntos] == ["2025-03-01", "2025-12-01", "2026-06-01"]
    assert [p["tipo"] for p in puntos] == ["acumulado", "anual", "acumulado"]
    assert puntos[1]["valor"] == round((32.8 / 32.0 - 1) * 100, 1) == 2.5


def test_un_acumulado_a_diciembre_cede_ante_el_anual_definitivo():
    ed = gestion.srt_parsear_provisorio(TEXTO_2026)
    ed = {**ed, "anio": 2026, "hasta_mes": 12, "periodo_txt": "enero a diciembre de 2026"}
    store = {"anuales": {"2025": {"casos": 1, "incidencia_mil": 10.0},
                         "2026": {"casos": 1, "incidencia_mil": 11.0}}, "ediciones": {}}
    gestion.srt_incorporar_edicion(store, ed, leida="2027-02-01")
    puntos = gestion.siniestralidad_puntos(store)
    assert [(p["fecha"], p["tipo"], p["valor"]) for p in puntos] == [("2026-12-01", "anual", 10.0)]


def test_releer_la_misma_edicion_no_cambia_el_store():
    store = _store_con(TEXTO_2026)
    antes = json.dumps(store, sort_keys=True)
    gestion.srt_incorporar_edicion(store, gestion.srt_parsear_provisorio(TEXTO_2026), leida="2027-01-01")
    assert json.dumps(store, sort_keys=True) == antes      # ni la fecha de lectura


def test_lee_la_serie_anual_del_excel():
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for hoja, titulo, fila in (
            ("Cuadro 2", "Cuadro 2: Casos notificados de AT y EP con días con baja o con secuelas incapacitantes",
             [332708, 326004]),
            ("Cuadro 4", "Cuadro 4: Indice de incidencia de AT y EP (por mil) según sector económico",
             [32.849388336782745, 32.01894005104061])):
        ws = wb.create_sheet(hoja)
        ws.append(["Serie histórica. Indicadores de AT y EP"])
        ws.append([titulo])
        ws.append(["Sector económico", 2025, 2024])
        ws.append(["Industria manufacturera", 1, 1])
        ws.append(["Total sistema", *fila])
    buf = io.BytesIO()
    wb.save(buf)
    # Los años van en orden creciente en la planilla real; acá van al revés a propósito.
    anuales = gestion.srt_parsear_serie_xlsx(buf.getvalue())
    assert set(anuales) == {"2024", "2025"}
    assert anuales["2025"]["casos"] == 332708
    assert anuales["2025"]["incidencia_mil"] == pytest.approx(32.849388336782745)


def test_la_serie_anual_falla_si_la_hoja_ya_no_es_la_que_era():
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for hoja in ("Cuadro 2", "Cuadro 4"):
        ws = wb.create_sheet(hoja)
        ws.append(["Serie histórica"])
        ws.append(["Cuadro: otra cosa"])
        ws.append(["Sector económico", 2024, 2025])
        ws.append(["Total sistema", 1, 2])
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(ValueError):
        gestion.srt_parsear_serie_xlsx(buf.getvalue())


# ── La tarjeta y la serie salen de la misma cuenta ───────────────────────────

@pytest.fixture
def store_tmp(tmp_path, monkeypatch):
    destino = tmp_path / "srt_siniestralidad.json"
    shutil.copy(STORE_REAL, destino)
    monkeypatch.setattr(gestion, "SRT_SINIESTRALIDAD_PATH", destino)
    return destino


def test_la_tarjeta_publica_el_ultimo_punto_del_store(store_tmp, monkeypatch):
    monkeypatch.setattr(gestion, "_srt_leer_provisorio",
                        lambda: gestion.srt_parsear_provisorio(TEXTO_2026))
    monkeypatch.setattr(gestion, "_http_get_resiliente",
                        lambda url: (_ for _ in ()).throw(RuntimeError("sin red")))
    antes = store_tmp.read_text(encoding="utf-8")
    card = gestion.fetch_siniestralidad_laboral()
    assert card["valor"] == -5.0
    assert card["fecha_dato"] == "2026-06-01"
    assert card["desactualizado"] is False
    assert "149.066" in card["detalle_txt"] and "159.671" in card["detalle_txt"]
    assert store_tmp.read_text(encoding="utf-8") == antes     # idempotente


def test_la_tarjeta_cae_a_cache_si_el_pdf_no_se_entiende(store_tmp, monkeypatch):
    def roto():
        raise ValueError("SRT: el PDF cambió de diseño")
    monkeypatch.setattr(gestion, "_srt_leer_provisorio", roto)
    assert gestion.fetch_siniestralidad_laboral() is None


def test_la_serie_termina_en_el_valor_de_la_tarjeta(store_tmp, monkeypatch):
    import descargar_series
    serie = descargar_series.fetch_siniestralidad_serie()
    meses = [f for f, _ in serie]
    assert meses[0] == "2020-12-01"
    assert meses == sorted(set(meses))                      # un valor por mes, sin saltos
    d = dict(serie)
    assert d["2021-11-01"] == -31.5 and d["2021-12-01"] == 32.4      # escalonada: cada año vale 12 meses
    assert d["2026-05-01"] == 2.6                            # el anual 2025 hasta que sale el provisorio
    assert d["2026-06-01"] == -5.0
    p = gestion.siniestralidad_puntos(gestion.srt_cargar_store())[-1]
    assert serie[-1][1] == p["valor"] == -5.0


def test_el_store_versionado_tiene_la_edicion_vigente_y_los_anuales():
    store = json.loads(STORE_REAL.read_text(encoding="utf-8"))
    assert "2026-06" in store["ediciones"]
    assert sorted(store["anuales"]) == ["2019", "2020", "2021", "2022", "2023", "2024", "2025"]
    puntos = gestion.siniestralidad_puntos(store)
    assert [p["fecha"][:4] for p in puntos] == ["2020", "2021", "2022", "2023", "2024", "2025", "2026"]
    assert puntos[-1]["valor"] == -5.0


# ── En el índice ─────────────────────────────────────────────────────────────

def test_reforma_laboral_pesa_50_25_25_con_la_siniestralidad():
    ind = itcg.DIMENSIONES_ITCG["reforma_laboral"]["indicadores"]
    assert ind == {"fal_modernizacion_laboral": 0.50, "siniestralidad_laboral": 0.25,
                   "juicios_trabajo_capital": 0.25}
    assert sum(ind.values()) == 1.0


def test_litigiosidad_queda_suspendida_pero_se_sigue_calculando():
    assert itcg.INDICADORES_SUSPENDIDOS["litigiosidad_laboral"]["adr"] == "0385"
    assert itcg.INDICADORES_SUSPENDIDOS["litigiosidad_laboral"]["dimension"] == "reforma_laboral"
    assert "litigiosidad_laboral" in gestion.INDICADORES_ESPERADOS
    assert "litigiosidad_laboral" in itcg.BANDAS_ITCG
    assert "litigiosidad_laboral" not in itcg.DIMENSIONES_ITCG["reforma_laboral"]["indicadores"]


def test_la_escala_es_la_convencion_de_los_juicios_y_esta_declarada():
    b = itcg.BANDAS_ITCG
    assert b["siniestralidad_laboral"] == b["litigiosidad_laboral"] == b["juicios_trabajo_capital"]
    assert procedencia_anclas.PROCEDENCIA["siniestralidad_laboral"][0] == "convencion"
    assert "litigiosidad_laboral" not in procedencia_anclas.PROCEDENCIA
    assert config.MAX_DIAS["siniestralidad_laboral"] >= 300      # un trimestre de margen


def test_el_corte_de_cinco_puntos_da_75():
    assert parametrica.puntaje_interpolado(-5.0, itcg.BANDAS_ITCG["siniestralidad_laboral"]) == 75.0


def test_la_dimension_puntua_con_la_siniestralidad_y_ignora_la_litigiosidad():
    base = {"fal_modernizacion_laboral": 50.0, "juicios_trabajo_capital": 12.4}
    con = itcg.calcular_itcg({**base, "siniestralidad_laboral": -5.0})
    dim = con["dimensiones"]["reforma_laboral"]
    # 0,50·55 + 0,25·75 + 0,25·40,2
    assert dim["puntaje"] == round(0.50 * 55 + 0.25 * 75.0 + 0.25 * 40.2, 1)
    # Un valor de litigiosidad, aunque llegue, no mueve nada.
    otro = itcg.calcular_itcg({**base, "siniestralidad_laboral": -5.0, "litigiosidad_laboral": 90.0})
    assert otro["dimensiones"]["reforma_laboral"]["puntaje"] == dim["puntaje"]
