"""La descarga del agregado ACLED tolera filas de relleno sin fecha."""
import sys
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
import gestion


def test_fila_de_relleno_sin_fecha_no_tumba_la_descarga(tmp_path, monkeypatch):
    monkeypatch.setattr(gestion, '_acled_credenciales', lambda: ('u', 'p'))
    monkeypatch.setattr(gestion, 'PROTESTAS_STORE_PATH', tmp_path / 'protestas.json')
    monkeypatch.setattr(gestion, '_PROTESTAS_STORE_MEMO', None)

    class Resp:
        text = ('<input name="form_build_id" value="f1">'
                '<a href="https://acleddata.com/system/files/2026/aggregated_data_latam.xlsx">x</a>')
        content = b'xlsx'

    class Ses:
        headers = {}
        def get(self, *a, **k): return Resp()
        def post(self, *a, **k): return Resp()
    monkeypatch.setattr(gestion.requests, 'Session', Ses)

    filas = [('WEEK', 'REGION', 'COUNTRY', 'ADMIN1', 'EVENT_TYPE', 'SUB_EVENT_TYPE', 'EVENTS'),
             (datetime(2026, 8, 22), 'South America', 'Argentina', 'Ciudad Autonoma de Buenos Aires',
              'Protests', 'Peaceful protest', 3),
             ('Source: ACLED', None, None, None, None, None, None),
             (datetime(2026, 8, 29), 'South America', 'Brazil', 'Sao Paulo', 'Protests', 'Peaceful protest', 5)]

    class Ws:
        def reset_dimensions(self): pass
        def iter_rows(self, values_only=True): return iter(filas)

    class Wb:
        sheetnames = ['Sheet1']
        def __getitem__(self, k): return Ws()
    import openpyxl
    monkeypatch.setattr(openpyxl, 'load_workbook', lambda *a, **k: Wb())

    store = gestion.actualizar_protestas_caba()
    assert store is not None
    assert store['_meta']['hasta_semana'] == '2026-08-29'
    assert store['mensual'] == {'2026-08': 3}
    assert store['mensual_nacional'] == {'2026-08': 3}
