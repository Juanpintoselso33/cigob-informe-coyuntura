"""Informalidad laboral SIN aportes propios (ADR-0413).

Mide el porcentaje de asalariados sin descuento jubilatorio que tampoco se
paga sus aportes por su cuenta (monotributo, autónomo): la cifra que el INDEC
llama «informal sin aportes propios». El asalariado que factura con
monotributo deja de contar como informal.

Fuentes, las dos planillas oficiales del INDEC (EPH, 31 aglomerados):

1. **Planilla de informalidad** `cuadros_informalidad_laboral_eph_MM_AA.xlsx`,
   «Cuadro 4.1». Trae la cifra ya calculada, desde el 4T-2023, pero se publica
   UNA vez al año (abril, con el 4T). Es la fuente oficial: manda donde existe.
2. **Planilla del informe de mercado de trabajo**
   `cuadros_tasas_indicadores_eph_MM_AA.xls` (MM = 03, 06, 09, 12: el mes en
   que sale cada trimestre), «Cuadro 1.5». Trae, por trimestre, el porcentaje
   de asalariados sin descuento jubilatorio y, dentro de ellos, cuántos tienen
   y cuántos no tienen aportes propios. Para los trimestres que la planilla
   anual todavía no cubre: sin_aportes = informales × proporción_sin_aportes.

El rezago se midió en la fuente: el 2T-2026 salió el 17-sep-2026, 169 días
después del inicio del trimestre (la fecha de cada punto es el inicio del
trimestre, como en las demás series EPH).

Los dos cálculos no coinciden al decimal (la planilla anual deja afuera a los
informales sin dato de aportes, y eso cambia el denominador: 4T-2025 da 30,3 en
la planilla y 30,6 por la cuenta). Por eso la planilla manda y la cuenta sólo
rellena lo posterior; cuando sale la planilla del año siguiente, reemplaza.

Guarda contra una errata de la fuente: la planilla trimestral trae, en el
2T-2025, 62,3 % «con aportes propios» (es el porcentaje de formales copiado en
la fila equivocada). Una proporción con aportes propios fuera de 3-40 % o que
no suma 100 se descarta en vez de entrar a la serie.
"""
from __future__ import annotations

import io
import logging
import re
from datetime import date
from functools import lru_cache

import requests

logger = logging.getLogger(__name__)

BASE_URL = "https://www.indec.gob.ar/ftp/cuadros/sociedad/"
PLANILLA_ANUAL = "cuadros_informalidad_laboral_eph_{mm:02d}_{yy:02d}.xlsx"
PLANILLA_TRIMESTRAL = "cuadros_tasas_indicadores_eph_{mm:02d}_{yy:02d}.xls"
MESES_TRIMESTRAL = (12, 9, 6, 3)
DESDE = "2023-10-01"          # 4T-2023: primer trimestre con el cuestionario nuevo
_UA = {"User-Agent": "Mozilla/5.0"}
_TRIM = {"1": "01", "2": "04", "3": "07", "4": "10"}


# ── descarga ────────────────────────────────────────────────────────────────
def _descargar(url: str) -> bytes | None:
    """Bytes del archivo o None. El INDEC contesta 200 con una página HTML
    cuando el archivo no existe: se distingue por la firma, no por el status."""
    try:
        r = requests.get(url, headers=_UA, timeout=60)
    except requests.RequestException as e:
        logger.warning("informalidad: no se pudo bajar %s (%s)", url, e)
        return None
    if r.status_code != 200 or not (r.content[:2] == b"PK" or r.content[:4] == b"\xd0\xcf\x11\xe0"):
        return None
    return r.content


def _candidatos(plantilla: str, meses, hoy: date, anios_atras: int):
    """URLs de la más nueva a la más vieja, sin fechas futuras."""
    for anio in range(hoy.year, hoy.year - anios_atras - 1, -1):
        for mm in meses:
            if (anio, mm) <= (hoy.year, hoy.month):
                yield BASE_URL + plantilla.format(mm=mm, yy=anio % 100)


# ── lectura de planillas a lista de filas ───────────────────────────────────
def _filas_xlsx(datos: bytes, prefijo_hoja: str) -> list:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(datos), data_only=True, read_only=True)
    for nombre in wb.sheetnames:
        if nombre.strip().startswith(prefijo_hoja):
            return [list(r) for r in wb[nombre].iter_rows(values_only=True)]
    raise ValueError(f"falta la hoja {prefijo_hoja!r}")


def _filas_xls(datos: bytes, prefijo_hoja: str) -> list:
    import xlrd
    wb = xlrd.open_workbook(file_contents=datos)
    for nombre in wb.sheet_names():
        if nombre.strip().startswith(prefijo_hoja):
            hoja = wb.sheet_by_name(nombre)
            return [hoja.row_values(i) for i in range(hoja.nrows)]
    raise ValueError(f"falta la hoja {prefijo_hoja!r}")


# ── parseo (puro, probado con filas sintéticas) ─────────────────────────────
def _fechas_por_columna(filas: list) -> dict:
    """{índice de columna: 'AAAA-MM-01'} desde las filas «Año 2023» y
    «4° trimestre» que encabezan todas las planillas EPH. El año de cada
    columna es el del último «Año» a su izquierda."""
    for i, fila in enumerate(filas[:15]):
        anios = {}
        for c, v in enumerate(fila):
            m = re.match(r"\s*A[ñn]o\s+(\d{4})", str(v or ""))
            if m:
                anios[c] = m.group(1)
        if not anios or i + 1 >= len(filas):
            continue
        out = {}
        for c, v in enumerate(filas[i + 1]):
            m = re.match(r"\s*([1-4])\s*[°º]?\s*trimestre", str(v or ""))
            izq = [k for k in anios if k <= c]
            if m and izq:
                out[c] = f"{anios[max(izq)]}-{_TRIM[m.group(1)]}-01"
        if out:
            return out
    raise ValueError("no se encontró el encabezado de trimestres")


def _fila(filas: list, etiqueta: str, desde: int = 0) -> tuple:
    """(índice, fila) de la primera fila cuya primera celda, sin espacios y en
    minúsculas, empieza con `etiqueta`."""
    for i in range(desde, len(filas)):
        if str(filas[i][0] or "").strip().lower().startswith(etiqueta):
            return i, filas[i]
    raise ValueError(f"falta la fila {etiqueta!r}")


def _num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def parsear_anual(filas: list) -> dict:
    """Cuadro 4.1: {fecha: % de asalariados informales sin aportes propios}."""
    cols = _fechas_por_columna(filas)
    _, sin = _fila(filas, "sin aportes propios")
    return {f: _num(sin[c]) for c, f in cols.items() if _num(sin[c]) is not None}


def parsear_trimestral(filas: list) -> dict:
    """Cuadro 1.5: {fecha: informales × proporción sin aportes propios}."""
    cols = _fechas_por_columna(filas)
    _, informal = _fila(filas, "sin descuento jubilatorio (informal)")
    i_con, con = _fila(filas, "con aportes propios")
    sin = filas[i_con + 1]
    if not str(sin[0]).strip().lower().startswith("sin aportes propios"):
        raise ValueError("la fila «sin aportes propios» no sigue a «con aportes propios»")
    out = {}
    for c, f in cols.items():
        inf, pc, ps = _num(informal[c]), _num(con[c]), _num(sin[c])
        if None in (inf, pc, ps):
            continue
        if not (3 <= pc <= 40) or abs(pc + ps - 100) > 0.35:
            logger.warning("informalidad: %s descartado, proporción con/sin aportes "
                           "propios inverosímil (%s / %s)", f, pc, ps)
            continue
        out[f] = round(inf * ps / 100, 1)
    return out


def combinar(anual: dict, trimestral: dict) -> list:
    """Serie ascendente [[fecha, %]] desde el 4T-2023. La planilla anual manda;
    la cuenta trimestral sólo aporta los trimestres posteriores al último que
    la anual cubre (empalmar en el medio mezclaría dos denominadores)."""
    ultimo_oficial = max(anual) if anual else ""
    todo = dict(anual)
    todo.update({f: v for f, v in trimestral.items() if f > ultimo_oficial})
    return [[f, v] for f, v in sorted(todo.items()) if f >= DESDE]


# ── API ─────────────────────────────────────────────────────────────────────
def _primero(plantilla, meses, anios_atras, hoy, leer, parsear):
    for url in _candidatos(plantilla, meses, hoy, anios_atras):
        datos = _descargar(url)
        if datos is None:
            continue
        try:
            return parsear(leer(datos)), url
        except Exception as e:  # una planilla con otro formato no tumba la otra
            logger.warning("informalidad: no se pudo leer %s (%s)", url, e)
    return {}, None


@lru_cache(maxsize=1)
def serie() -> list:
    """[[fecha, %]] ascendente desde el 4T-2023; falla en voz alta si no hay nada."""
    hoy = date.today()
    anual, u1 = _primero(PLANILLA_ANUAL, range(12, 0, -1), 1, hoy,
                         lambda d: _filas_xlsx(d, "Cuadro 4.1"), parsear_anual)
    trim, u2 = _primero(PLANILLA_TRIMESTRAL, MESES_TRIMESTRAL, 1, hoy,
                        lambda d: _filas_xls(d, "Cuadro 1.5"), parsear_trimestral)
    out = combinar(anual, trim)
    if not out:
        raise RuntimeError("informalidad sin aportes propios: ninguna planilla del INDEC legible")
    # sin el 4T-2023 la serie no tiene base y el componente queda sin puntaje en
    # silencio: mejor fallar y que el pipeline conserve lo anterior
    if out[0][0] != DESDE:
        raise RuntimeError(f"informalidad sin aportes propios: falta la base {DESDE} "
                           f"(planilla anual: {u1 or 'ninguna legible'})")
    logger.info("informalidad sin aportes propios: %d trimestres (%s | %s)", len(out), u1, u2)
    return out


def ultimo() -> dict:
    """{'valor': proporción, 'fecha': inicio del trimestre}: mismo formato que
    el resto de los colectores INDEC (publicar.py la pasa a % con _red)."""
    fecha, pct = serie()[-1]
    return {"valor": round(pct / 100, 4), "fecha": fecha}
